#!/usr/bin/env python3
"""Compila un template Atlas (.docx o .pptx) sostituendo i segnaposto tra parentesi quadre,
riscrivendo le tabelle e togliendo i paragrafi di istruzione. Conserva stili, font e layout
del template: il documento resta con l'identità visiva Atlas.

Uso:
  python fill_template.py TEMPLATE OUTPUT SPEC.json

SPEC.json (tutte le chiavi sono facoltative):
{
  "replace": {"[Cliente]": "Confindustria Como", "[Data]": "2026-10-05"},
  "tables": [
    {"header": "Processo | Area",            # inizio della riga di intestazione (celle unite da " | ")
     "rows": [["Onboarding associati", "Servizi", "40/mese", "..."], ...],
     "slide": 6}                              # solo pptx, facoltativo: limita la ricerca a una slide
  ],
  "remove_paragraphs": ["Mantenere una sola delle tre varianti"],   # sottostringhe
  "remove_boxes": ["Regole di compilazione", "Criterio di uscita"],  # solo docx: riquadri a una cella
                                                                       # il cui testo inizia così
  "boxes": {"Contesto nel programma Atlas": [                          # solo docx: riscrive il riquadro
      "**Da dove veniamo.** ...", "**Dove siamo.** ...", "**Dove andiamo.** ..."]},  # tenendo titolo e stile
  "delete_slides": [5],                                                # solo pptx, numeri 1-based
  "charts": [                                                          # solo pptx
    {"slide": 5, "categories": ["Strategia e sponsorship", "Processi", "Dati e integrazioni",
                                "Competenze", "Governance e rischio", "Misurazione"],
     "series": {"Oggi": [2, 2, 1, 1, 1, 1], "Target": [4, 3, 3, 3, 3, 3]}},
    {"slide": 8, "xy": {"Processi": [[4.2, 4.3], [3.6, 3.2]]}}          # [fattibilità, valore]
  ]
}

Regole:
- "replace" lavora prima sul singolo run (i segnaposto dei template sono run in corsivo:
  il testo inserito perde il corsivo), poi sul paragrafo intero se il segnaposto è spezzato.
- "tables" sostituisce le righe dati (dalla seconda in poi) usando la prima riga dati come
  modello: righe in più vengono clonate, righe in meno eliminate. Le celle vuote restano vuote.
- A fine esecuzione stampa i segnaposto rimasti: quelli voluti (es. [importo], che compila il
  Calculator) vanno bene, gli altri vanno compilati o dichiarati come punti aperti.
"""
import copy, json, re, sys

PH = re.compile(r"\[[^\]\n]{2,90}\]")


def _replace_in_runs(runs, mapping):
    changed = False
    for r in runs:
        t = r.text
        new = t
        for k, v in mapping.items():
            if k in new:
                new = new.replace(k, str(v))
        if new != t:
            r.text = new
            try:
                if r.font.italic and not PH.search(new):
                    r.font.italic = False
            except AttributeError:
                pass
            changed = True
    return changed


def _replace_in_paragraph(p, mapping):
    _replace_in_runs(p.runs, mapping)
    full = "".join(r.text for r in p.runs)
    if any(k in full for k in mapping):  # segnaposto spezzato su più run
        new = full
        for k, v in mapping.items():
            new = new.replace(k, str(v))
        runs = p.runs
        if runs:
            runs[0].text = new
            try:
                runs[0].font.italic = False if not PH.search(new) else runs[0].font.italic
            except AttributeError:
                pass
            for r in runs[1:]:
                r.text = ""


def _set_cell_text(cell, text):
    paras = cell.paragraphs if hasattr(cell, "paragraphs") else cell.text_frame.paragraphs
    first = paras[0]
    runs = first.runs
    if runs:
        runs[0].text = str(text)
        try:
            runs[0].font.italic = False
        except AttributeError:
            pass
        for r in runs[1:]:
            r.text = ""
    else:
        first.add_run().text = str(text) if hasattr(first, "add_run") else None
    for extra in paras[1:]:
        el = extra._p if hasattr(extra, "_p") else extra._element
        el.getparent().remove(el)


# ---------------------------------------------------------------- DOCX
def fill_docx(src, dst, spec):
    import docx
    from docx.table import _Cell
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document(src)
    mapping = spec.get("replace", {})
    removes = spec.get("remove_paragraphs", [])

    def all_paragraphs():
        for p in d.paragraphs:
            yield p
        for t in iter_tables():
            for tr in t._tbl.tr_lst:
                for tc in tr.tc_lst:
                    for p in _Cell(tc, t).paragraphs:
                        yield p
        for s in d.sections:
            for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                for p in part.paragraphs:
                    yield p

    def iter_tables():
        stack = list(d.tables)
        while stack:
            t = stack.pop(0)
            yield t
            for tr in t._tbl.tr_lst:
                for tc in tr.tc_lst:
                    stack.extend(_Cell(tc, t).tables)

    def row_cells(t, tr):
        return [_Cell(tc, t) for tc in tr.tc_lst]

    # riquadri di istruzione del template (tabelle a una cella): "remove_boxes": ["Regole di compilazione", ...]
    for t in list(d.tables):
        first = t._tbl.tr_lst[0].tc_lst[0] if t._tbl.tr_lst and t._tbl.tr_lst[0].tc_lst else None
        if first is not None and any(_Cell(first, t).text.strip().startswith(k) for k in spec.get("remove_boxes", [])):
            t._tbl.getparent().remove(t._tbl)

    # riquadri da riscrivere con il contesto del cliente: "boxes": {"Contesto nel programma Atlas": ["**Da dove veniamo.** ...", ...]}
    for key, new_lines in spec.get("boxes", {}).items():
        for t in d.tables:
            trs = t._tbl.tr_lst
            if not trs or not trs[0].tc_lst:
                continue
            tc = trs[0].tc_lst[0]
            cell = _Cell(tc, t)
            if not cell.text.strip().startswith(key):
                continue
            ps = cell.paragraphs
            proto = ps[1]._p if len(ps) > 1 else ps[0]._p
            proto_runs = proto.findall(qn("w:r"))
            for p in ps[1:]:
                p._p.getparent().remove(p._p)
            for line in new_lines:
                np_ = copy.deepcopy(proto)
                for r in np_.findall(qn("w:r")):
                    np_.remove(r)
                m = re.match(r"\*\*(.+?)\*\*\s*(.*)", line)
                parts = [(m.group(1) + " ", True), (m.group(2), False)] if m else [(line, False)]
                for txt, bold in parts:
                    if not txt:
                        continue
                    src = proto_runs[0] if (bold and proto_runs) else (proto_runs[-1] if proto_runs else None)
                    r = copy.deepcopy(src) if src is not None else OxmlElement("w:r")
                    for tnode in r.findall(qn("w:t")):
                        r.remove(tnode)
                    rp = r.find(qn("w:rPr"))
                    if rp is not None:
                        for tag in ("w:b", "w:bCs", "w:i", "w:iCs"):
                            for e in rp.findall(qn(tag)):
                                rp.remove(e)
                        if bold:
                            rp.insert(1, OxmlElement("w:b"))
                    tn = OxmlElement("w:t"); tn.set(qn("xml:space"), "preserve"); tn.text = txt
                    r.append(tn)
                    np_.append(r)
                tc.append(np_)
            break
        else:
            print(f"[AVVISO] riquadro non trovato: {key}")

    for tspec in spec.get("tables", []):
        hdr = tspec["header"].strip()
        target = None
        for t in iter_tables():
            if not t._tbl.tr_lst:
                continue
            h = " | ".join(c.text.strip() for c in row_cells(t, t._tbl.tr_lst[0]))
            if h.startswith(hdr):
                target = t
                break
        if target is None:
            print(f"[AVVISO] tabella non trovata: {hdr}")
            continue
        trs = target._tbl.tr_lst
        if len(trs) < 2:
            print(f"[AVVISO] tabella senza righe dati: {hdr}")
            continue
        proto = copy.deepcopy(trs[1])
        for tr in trs[1:]:
            target._tbl.remove(tr)
        for row in tspec["rows"]:
            tr = copy.deepcopy(proto)
            target._tbl.append(tr)
            cells = row_cells(target, tr)
            for i, c in enumerate(cells):
                _set_cell_text(c, row[i] if i < len(row) else "")

    for p in list(all_paragraphs()):
        txt = p.text
        if any(s in txt for s in removes):
            p._p.getparent().remove(p._p)
            continue
        if mapping:
            _replace_in_paragraph(p, mapping)

    d.save(dst)
    leftovers = []
    for p in all_paragraphs():
        leftovers += PH.findall(p.text)
    return leftovers


# ---------------------------------------------------------------- PPTX
def fill_pptx(src, dst, spec):
    from pptx import Presentation

    prs = Presentation(src)
    mapping = spec.get("replace", {})
    removes = spec.get("remove_paragraphs", [])

    for tspec in spec.get("tables", []):
        hdr = tspec["header"].strip()
        found = False
        for i, s in enumerate(prs.slides, 1):
            if tspec.get("slide") and tspec["slide"] != i:
                continue
            for sh in s.shapes:
                if not getattr(sh, "has_table", False) or not sh.has_table:
                    continue
                tbl = sh.table
                h = " | ".join(c.text.strip() for c in tbl.rows[0].cells)
                if not h.startswith(hdr):
                    continue
                found = True
                xtbl = tbl._tbl
                trs = xtbl.tr_lst
                proto = copy.deepcopy(trs[1])
                for tr in trs[1:]:
                    xtbl.remove(tr)
                for row in tspec["rows"]:
                    tr = copy.deepcopy(proto)
                    xtbl.append(tr)
                from pptx.table import _Row
                for ri, row in enumerate(tspec["rows"], 1):
                    r = tbl.rows[ri]
                    for ci, c in enumerate(r.cells):
                        _set_cell_text(c, row[ci] if ci < len(row) else "")
                break
            if found:
                break
        if not found:
            print(f"[AVVISO] tabella non trovata: {hdr}")

    def paragraphs():
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        yield p
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        for c in r.cells:
                            for p in c.text_frame.paragraphs:
                                yield p

    for p in list(paragraphs()):
        txt = "".join(r.text for r in p.runs)
        if any(s in txt for s in removes):
            for r in p.runs:
                r.text = ""
            continue
        if mapping:
            _replace_in_paragraph(p, mapping)

    from pptx.chart.data import CategoryChartData, XyChartData
    for cspec in spec.get("charts", []):
        s = prs.slides[cspec["slide"] - 1]
        charts = [sh for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
        if not charts:
            print(f"[AVVISO] nessun grafico nella slide {cspec['slide']}")
            continue
        ch = charts[cspec.get("index", 0)].chart
        if "xy" in cspec:  # dispersione: {"Serie": [[x, y], ...]}
            cd = XyChartData()
            for name, pts in cspec["xy"].items():
                ser = cd.add_series(name)
                for x, y in pts:
                    ser.add_data_point(x, y)
        else:  # categorie: {"categories": [...], "series": {"Oggi": [...], "Target": [...]}}
            cd = CategoryChartData()
            cd.categories = cspec["categories"]
            for name, vals in cspec["series"].items():
                cd.add_series(name, vals)
        ch.replace_data(cd)

    for n in sorted(spec.get("delete_slides", []), reverse=True):
        sldIdLst = prs.slides._sldIdLst
        sid = sldIdLst[n - 1]
        prs.part.drop_rel(sid.rId)
        sldIdLst.remove(sid)

    prs.save(dst)
    left = []
    for p in paragraphs():
        left += PH.findall("".join(r.text for r in p.runs))
    return left


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    src, dst, spec_path = sys.argv[1:]
    spec = json.load(open(spec_path, encoding="utf-8"))
    if src.lower().endswith(".docx"):
        left = fill_docx(src, dst, spec)
    elif src.lower().endswith(".pptx"):
        left = fill_pptx(src, dst, spec)
    else:
        sys.exit("Formato non supportato: usa .docx o .pptx")
    print(f"Salvato: {dst}")
    if left:
        from collections import Counter
        print("Segnaposto rimasti (compilarli o dichiararli come punti aperti):")
        for k, n in Counter(left).most_common():
            print(f"  {n}× {k}")
    else:
        print("Nessun segnaposto rimasto.")


if __name__ == "__main__":
    main()
